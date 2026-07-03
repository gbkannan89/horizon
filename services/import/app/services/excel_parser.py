import time
from decimal import Decimal
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import load_workbook

from app.models.schemas import ParsedTransaction
from app.services.csv_parser import (
    BANK_FORMATS,
    build_generic_mapping,
    detect_event_type,
    detect_format,
    parse_amount,
    parse_date,
)


def _detect_header_row(rows: list[list]) -> int | None:
    for i, row in enumerate(rows):
        header_text = " ".join(str(c).strip().lower() for c in row if c is not None)
        kw = ("amount", "description", "narration", "debit", "credit")
        if "date" in header_text and any(w in header_text for w in kw):
            return i
    return None


def _find_data_sheet(wb) -> str | None:
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row < 2:
            continue
        sample = " ".join(str(c).lower() for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True)) if c)
        kw = ("amount", "description", "debit", "credit", "narration", "balance")
        if "date" in sample and any(w in sample for w in kw):
            return sheet_name
    return wb.sheetnames[0] if wb.sheetnames else None


def parse_excel(file_content: bytes | str | Path) -> list[ParsedTransaction]:
    if isinstance(file_content, str):
        file_content = file_content.encode("utf-8")
    elif isinstance(file_content, Path):
        file_content = file_content.read_bytes()

    with NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        tmp.write(file_content)
        tmp_path = tmp.name

    wb = None
    try:
        wb = load_workbook(tmp_path, read_only=True, data_only=True)
        sheet_name = _find_data_sheet(wb)
        if sheet_name is None:
            return []

        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows or all(c is None for c in rows[0]):
            return []

        header_row_idx = _detect_header_row(rows)
        if header_row_idx is None:
            header_row_idx = 0

        if header_row_idx >= len(rows):
            return []

        headers = [str(c) if c is not None else "" for c in rows[header_row_idx]]
        data_rows = rows[header_row_idx + 1:]

        fmt = detect_format(headers)
        if fmt == "generic":
            mapping = build_generic_mapping(headers)
        else:
            mapping = BANK_FORMATS[fmt]

        transactions: list[ParsedTransaction] = []
        for row_num, row in enumerate(data_rows, start=header_row_idx + 2):
            row_dict = {headers[i]: str(v) if v is not None else "" for i, v in enumerate(row)}
            if all(v == "" for v in row_dict.values()):
                continue
            txn = _parse_excel_row(row_dict, mapping)
            if txn is not None:
                transactions.append(txn)

        return transactions

    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass
        for _ in range(5):
            try:
                Path(tmp_path).unlink()
                break
            except PermissionError:
                time.sleep(0.1)


def _parse_excel_row(row: dict[str, str], mapping) -> ParsedTransaction | None:
    date_val = row.get(mapping.date_col, "").strip()
    desc = row.get(mapping.description_col, "").strip()

    if not date_val and not desc:
        return None

    event_date = parse_date(date_val, mapping.date_format)
    effective_date = event_date

    debit_raw = row.get(mapping.debit_col, "").strip() if mapping.debit_col else ""
    credit_raw = row.get(mapping.credit_col, "").strip() if mapping.credit_col else ""
    amount_raw = row.get(mapping.amount_col, "").strip() if mapping.amount_col else ""

    debit = parse_amount(debit_raw) if debit_raw else Decimal("0")
    credit = parse_amount(credit_raw) if credit_raw else Decimal("0")
    amount = parse_amount(amount_raw) if amount_raw else Decimal("0")

    if amount == 0:
        if debit > 0:
            amount = -debit
        elif credit > 0:
            amount = credit

    ref = row.get(mapping.reference_col, "").strip() if mapping.reference_col else ""
    event_type = detect_event_type(desc, amount)

    return ParsedTransaction(
        event_type=event_type,
        amount=abs(amount),
        currency="INR",
        event_date=event_date,
        effective_date=effective_date,
        description=desc,
        reference=ref,
    )
